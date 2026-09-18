"""Consolidated services for Prompts 13-32: jobs, clients, vendors, templates,
applications (convert-to-lead), reports, action-required, lead-inbox, integrations,
import batches, settings lists, tags, global search."""
import base64
import csv
import hashlib
import io
from datetime import timedelta
from core.config import settings
from core.errors import AppError
from core.transactions import atomic
from models.entities import new_id, now_utc
from models.recruitment import ACTIVE_STATUSES, FUNNEL_ORDER, STATUS_LABELS
from services import scope_service, recruitment_service as rs

# ---------------- crypto for integration secrets ----------------
def _fernet():
    from cryptography.fernet import Fernet
    key = base64.urlsafe_b64encode(hashlib.sha256(settings.jwt_secret.encode()).digest())
    return Fernet(key)


def _encrypt(v):
    return _fernet().encrypt(v.encode()).decode()


def _mask(v):
    if not v:
        return ""
    return "••••" + v[-4:] if len(v) > 4 else "••••"


def clean(d):
    d = dict(d); d.pop("_id", None); return d


# ======================= JOBS =======================
async def list_jobs(db, actor, perms, *, search="", status=""):
    q = {"archived": {"$ne": True}}
    if status:
        q["status"] = status
    if search:
        q["$or"] = [{"title": {"$regex": search, "$options": "i"}},
                    {"client": {"$regex": search, "$options": "i"}},
                    {"location": {"$regex": search, "$options": "i"}}]
    rows = await db.jobs.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    for r in rows:
        r["linked_leads"] = await db.leads.count_documents({"job": r["title"]})
    return rows


async def save_job(db, actor, data, job_id=None):
    if not data.get("title") or not data.get("client"):
        raise AppError("bad_request", "Position title and client are required", 400)
    now = now_utc()
    if job_id:
        job = await db.jobs.find_one({"id": job_id})
        if not job:
            raise AppError("not_found", "Job not found", 404)
        data["updated_at"] = now
        await db.jobs.update_one({"id": job_id}, {"$set": data})
        return clean(await db.jobs.find_one({"id": job_id}))
    doc = {"id": new_id(), "status": "active", **data, "archived": False,
           "created_at": now, "updated_at": now, "created_by": actor["id"]}
    await db.jobs.insert_one(doc)
    return clean(doc)


async def archive_job(db, job_id):
    await db.jobs.update_one({"id": job_id}, {"$set": {"archived": True, "status": "closed"}})
    return {"archived": True}


# ======================= CLIENTS =======================
async def list_clients(db, actor, perms, *, search=""):
    q = {"archived": {"$ne": True}}
    if search:
        q["$or"] = [{"name": {"$regex": search, "$options": "i"}},
                    {"company": {"$regex": search, "$options": "i"}},
                    {"contact_person": {"$regex": search, "$options": "i"}},
                    {"contact_phone": {"$regex": search, "$options": "i"}}]
    rows = await db.clients.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    for r in rows:
        base = {"client": r["name"]}
        r["stats"] = {
            "submitted": await db.leads.count_documents(base),
            "interviewed": await db.leads.count_documents({**base, "status": {"$in": ["lineup", "selected", "joined"]}}),
            "selected": await db.leads.count_documents({**base, "status": {"$in": ["selected", "joined"]}}),
            "joined": await db.leads.count_documents({**base, "status": "joined"}),
        }
    return rows


async def save_client(db, actor, data, client_id=None):
    if not data.get("name"):
        raise AppError("bad_request", "Client name is required", 400)
    now = now_utc()
    if client_id:
        if not await db.clients.find_one({"id": client_id}):
            raise AppError("not_found", "Client not found", 404)
        data["updated_at"] = now
        await db.clients.update_one({"id": client_id}, {"$set": data})
        return clean(await db.clients.find_one({"id": client_id}))
    doc = {"id": new_id(), **data, "archived": False, "created_at": now, "updated_at": now, "created_by": actor["id"]}
    await db.clients.insert_one(doc)
    return clean(doc)


async def archive_client(db, client_id):
    await db.clients.update_one({"id": client_id}, {"$set": {"archived": True}})
    return {"archived": True}


# ======================= VENDORS =======================
VENDOR_STAGES = ["new", "contacted", "followup", "meeting", "proposal_sent", "negotiation", "empanelled", "rejected"]


async def list_vendors(db, actor, perms, *, search="", stage=""):
    ids = await scope_service.recruiter_scope_ids(db, actor, perms)
    q = {} if ids is None else {"owner_id": {"$in": ids}}
    if stage:
        q["stage"] = stage
    if search:
        q["$or"] = [{"company": {"$regex": search, "$options": "i"}},
                    {"contact_person": {"$regex": search, "$options": "i"}},
                    {"phone": {"$regex": search, "$options": "i"}}]
    rows = await db.vendors.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    return rows


async def vendor_summary(db, actor, perms):
    ids = await scope_service.recruiter_scope_ids(db, actor, perms)
    base = {} if ids is None else {"owner_id": {"$in": ids}}
    return {
        "empanelled": await db.vendors.count_documents({**base, "stage": "empanelled"}),
        "proposals": await db.vendors.count_documents({**base, "stage": {"$in": ["proposal_sent", "negotiation"]}}),
        "active": await db.vendors.count_documents({**base, "stage": {"$nin": ["empanelled", "rejected"]}}),
    }


async def save_vendor(db, actor, perms, data, vendor_id=None):
    if not data.get("company"):
        raise AppError("bad_request", "Company name is required", 400)
    now = now_utc()
    if vendor_id:
        v = await db.vendors.find_one({"id": vendor_id})
        if not v:
            raise AppError("not_found", "Vendor not found", 404)
        await scope_service.assert_lead_access(db, actor, perms, {"owner_id": v["owner_id"]})
        data["updated_at"] = now
        await db.vendors.update_one({"id": vendor_id}, {"$set": data})
        return clean(await db.vendors.find_one({"id": vendor_id}))
    doc = {"id": new_id(), "stage": data.get("stage") or "new", **data,
           "owner_id": actor["id"], "created_at": now, "updated_at": now, "created_by": actor["id"]}
    await db.vendors.insert_one(doc)
    return clean(doc)


async def delete_vendor(db, actor, perms, vendor_id):
    v = await db.vendors.find_one({"id": vendor_id})
    if not v:
        raise AppError("not_found", "Vendor not found", 404)
    await scope_service.assert_lead_access(db, actor, perms, {"owner_id": v["owner_id"]})
    await db.vendors.delete_one({"id": vendor_id})
    return {"deleted": True}


# ======================= TEMPLATES (unified) =======================
async def list_templates(db, *, channel="", search=""):
    q = {"archived": {"$ne": True}}
    if channel:
        q["channel"] = channel
    if search:
        q["$or"] = [{"name": {"$regex": search, "$options": "i"}},
                    {"category": {"$regex": search, "$options": "i"}},
                    {"body": {"$regex": search, "$options": "i"}}]
    return await db.templates.find(q, {"_id": 0}).sort("updated_at", -1).to_list(500)


async def save_template(db, actor, data, tpl_id=None):
    if not data.get("name") or not data.get("body"):
        raise AppError("bad_request", "Template name and body are required", 400)
    if data.get("channel") not in ("whatsapp", "email"):
        raise AppError("bad_request", "Channel must be whatsapp or email", 400)
    now = now_utc()
    if tpl_id:
        if not await db.templates.find_one({"id": tpl_id}):
            raise AppError("not_found", "Template not found", 404)
        data["updated_at"] = now; data["updated_by"] = actor.get("name")
        await db.templates.update_one({"id": tpl_id}, {"$set": data})
        return clean(await db.templates.find_one({"id": tpl_id}))
    doc = {"id": new_id(), "provider_approved": False, **data, "archived": False,
           "created_at": now, "updated_at": now, "updated_by": actor.get("name"), "created_by": actor["id"]}
    await db.templates.insert_one(doc)
    return clean(doc)


async def delete_template(db, tpl_id):
    await db.templates.update_one({"id": tpl_id}, {"$set": {"archived": True}})
    return {"deleted": True}


# ======================= APPLICATIONS =======================
async def list_applications(db, actor, perms, *, search="", status=""):
    q = {"archived": {"$ne": True}}
    if status and status != "all":
        q["status"] = status
    if search:
        q["$or"] = [{"name": {"$regex": search, "$options": "i"}},
                    {"phone": {"$regex": search, "$options": "i"}},
                    {"role": {"$regex": search, "$options": "i"}}]
    return await db.applications.find(q, {"_id": 0}).sort("applied_at", -1).to_list(500)


async def update_application_status(db, app_id, status):
    if status not in ("new", "shortlisted", "converted"):
        raise AppError("bad_request", "Invalid status", 400)
    await db.applications.update_one({"id": app_id}, {"$set": {"status": status}})
    return clean(await db.applications.find_one({"id": app_id}))


async def archive_application(db, app_id):
    await db.applications.update_one({"id": app_id}, {"$set": {"archived": True}})
    return {"archived": True}


async def convert_application(db, actor, perms, app_id, assign_to=None, create_followup=True):
    app = await db.applications.find_one({"id": app_id})
    if not app:
        raise AppError("not_found", "Application not found", 404)
    # idempotent: if already converted and linked, return the linked lead
    if app.get("status") == "converted" and app.get("lead_id"):
        lead = await db.leads.find_one({"id": app["lead_id"]}, {"_id": 0})
        if lead:
            return {"converted": False, "already": True, "lead": rs.clean(lead)}
    norm = rs._normalize_phone(app.get("phone", ""))
    existing = await db.leads.find_one({"phone_normalized": norm}) if norm else None
    if existing:
        await db.applications.update_one({"id": app_id},
            {"$set": {"status": "converted", "lead_id": existing["id"]}})
        return {"converted": False, "existing": True, "lead": rs.clean(existing)}

    owner = assign_to or actor["id"]
    now = now_utc()
    lead_id = new_id()
    doc = {
        "id": lead_id, "lead_code": rs._lead_code(), "name": app["name"], "phone": app["phone"],
        "phone_normalized": norm, "alt_phone": None, "email": app.get("email"),
        "city": app.get("city"), "age": None, "gender": None, "qualification": None,
        "experience": app.get("experience"), "current_salary": None, "expected_salary": None,
        "notice_period": None, "source": app.get("source") or "application",
        "role_applied": app.get("role"), "client": None, "job": app.get("role"),
        "priority": "medium", "status": "new", "status_changed_at": now, "owner_id": owner,
        "team_id": None, "tags": [], "notes": None, "last_call_outcome": None, "call_count": 0,
        "last_activity_at": now, "next_followup_at": None, "closure_reason": None,
        "expected_joining_at": None,
        "campaign": app.get("campaign"), "utm_source": app.get("utm_source"),
        "utm_medium": app.get("utm_medium"), "utm_campaign": app.get("utm_campaign"),
        "provider_lead_id": app.get("provider_lead_id"), "first_source": app.get("source"),
        "status_history": [{"status": "new", "at": now, "by": actor["id"]}],
        "created_at": now, "updated_at": now, "created_by": actor["id"], "updated_by": actor["id"],
    }
    async with atomic() as s:
        await db.leads.insert_one(doc, session=s)
        await db.applications.update_one({"id": app_id},
            {"$set": {"status": "converted", "lead_id": lead_id}}, session=s)
        await rs.write_activity(db, lead_id, "created", actor,
                                f"Converted from application ({doc['lead_code']})",
                                {"source": doc["source"], "campaign": doc.get("campaign")}, session=s)
        if create_followup:
            await db.followups.insert_one({"id": new_id(), "lead_id": lead_id, "lead_name": doc["name"],
                "recruiter_id": owner, "due_at": now + timedelta(days=1), "status": "pending",
                "reason": "First call", "notes": None, "created_at": now, "created_by": actor["id"]}, session=s)
            await db.leads.update_one({"id": lead_id},
                {"$set": {"next_followup_at": now + timedelta(days=1)}}, session=s)
    from services import notify_service
    await notify_service.notify(db, recipient_id=owner, type_="lead_assigned",
                                title="New lead from application", message=doc["name"],
                                entity_type="lead", entity_id=lead_id, action_url=f"/admin/leads?open={lead_id}")
    return {"converted": True, "lead": rs.clean(doc)}


# ======================= LEAD INBOX =======================
async def lead_inbox(db, actor, perms, *, search="", source=""):
    scope_q = await scope_service.leads_filter(db, actor, perms, None)
    # source summary
    pipeline = [{"$match": scope_q}, {"$group": {"_id": "$source",
        "total": {"$sum": 1},
        "connected": {"$sum": {"$cond": [{"$gt": ["$call_count", 0]}, 1, 0]}},
        "interview": {"$sum": {"$cond": [{"$in": ["$status", ["lineup", "selected", "joined"]]}, 1, 0]}},
        "selected": {"$sum": {"$cond": [{"$in": ["$status", ["selected", "joined"]]}, 1, 0]}},
        "joined": {"$sum": {"$cond": [{"$eq": ["$status", "joined"]}, 1, 0]}}}}]
    summary = [{"source": r["_id"] or "unknown", **{k: r[k] for k in ("total", "connected", "interview", "selected", "joined")}}
               async for r in db.leads.aggregate(pipeline)]
    q = dict(scope_q)
    if source and source != "all":
        q["source"] = source
    if search:
        q["$or"] = [{"name": {"$regex": search, "$options": "i"}},
                    {"phone": {"$regex": search, "$options": "i"}},
                    {"email": {"$regex": search, "$options": "i"}}]
    rows = await db.leads.find(q, {"_id": 0}).sort("created_at", -1).to_list(300)
    await rs._attach_flags(db, actor, perms, rows, scope_q)
    return {"summary": summary, "items": rows}


# ======================= REPORTS =======================
async def _recruiter_ids(db, actor, perms):
    ids = await scope_service.recruiter_scope_ids(db, actor, perms)
    users = await db.users.find({} if ids is None else {"id": {"$in": ids}},
                                {"_id": 0, "id": 1, "name": 1}).to_list(500)
    return {u["id"]: u["name"] for u in users}


async def reports(db, actor, perms, *, tab="leaderboard"):
    umap = await _recruiter_ids(db, actor, perms)
    rid = list(umap.keys())
    org = await db.organization_settings.find_one({"key": "org"}) or {}
    targets = org.get("targets", {"calls_per_day": 50, "connected_per_day": 20, "lineups_per_day": 3, "joinings_per_month": 5})
    lead_q = {"owner_id": {"$in": rid}}

    if tab == "funnel":
        out = []
        for st in FUNNEL_ORDER:
            out.append({"stage": st, "label": STATUS_LABELS.get(st, st),
                        "count": await db.leads.count_documents({**lead_q, "status": st})})
        return {"funnel": out}

    if tab == "aging":
        now = now_utc()
        buckets = {"new": 0, "1d": 0, "3d": 0, "7d": 0, "15d+": 0}
        async for l in db.leads.find({**lead_q, "status": {"$in": ACTIVE_STATUSES}}, {"_id": 0, "created_at": 1}):
            age = (now - rs._as_utc(l.get("created_at") or now)).days
            if age < 1:
                buckets["new"] += 1
            elif age < 3:
                buckets["1d"] += 1
            elif age < 7:
                buckets["3d"] += 1
            elif age < 15:
                buckets["7d"] += 1
            else:
                buckets["15d+"] += 1
        return {"aging": buckets}

    if tab == "missed":
        now = now_utc()
        rows = await db.followups.find({"recruiter_id": {"$in": rid}, "status": "pending",
                                        "due_at": {"$lt": now}}, {"_id": 0}).sort("due_at", 1).to_list(300)
        for r in rows:
            l = await db.leads.find_one({"id": r["lead_id"]}, {"_id": 0, "phone": 1, "priority": 1, "status": 1})
            r["recruiter"] = umap.get(r["recruiter_id"])
            r["phone"] = (l or {}).get("phone"); r["priority"] = (l or {}).get("priority")
            r["lead_status"] = (l or {}).get("status")
            r["delay_hours"] = round((now - rs._as_utc(r["due_at"])).total_seconds() / 3600, 1)
        return {"missed": rows}

    # leaderboard + targets share per-recruiter aggregation
    today = now_utc().replace(hour=0, minute=0, second=0, microsecond=0)
    month_start = now_utc().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    board = []
    for uid, name in umap.items():
        calls = await db.calls.count_documents({"recruiter_id": uid, "created_at": {"$gte": today}})
        connected = await db.calls.count_documents({"recruiter_id": uid, "connected": True, "created_at": {"$gte": today}})
        lineups = await db.leads.count_documents({"owner_id": uid, "status": "lineup"})
        attended = await db.interviews.count_documents({"recruiter_id": uid, "status": "attended"})
        selected = await db.leads.count_documents({"owner_id": uid, "status": "selected"})
        joined = await db.joinings.count_documents({"recruiter_id": uid, "status": "joined", "created_at": {"$gte": month_start}})
        missed = await db.followups.count_documents({"recruiter_id": uid, "status": "pending", "due_at": {"$lt": now_utc()}})
        score = calls * 1 + connected * 2 + lineups * 5 + selected * 8 + joined * 15 - missed * 2
        board.append({"recruiter": name, "recruiter_id": uid, "calls": calls, "connected": connected,
                      "lineups": lineups, "attended": attended, "selected": selected, "joined": joined,
                      "missed": missed, "score": score})
    board.sort(key=lambda x: x["score"], reverse=True)
    for i, b in enumerate(board):
        b["rank"] = i + 1
    if tab == "targets":
        cards = []
        for b in board:
            cards.append({"recruiter": b["recruiter"], "metrics": [
                {"metric": "Calls / day", "actual": b["calls"], "target": targets["calls_per_day"]},
                {"metric": "Connected / day", "actual": b["connected"], "target": targets["connected_per_day"]},
                {"metric": "Lineups / day", "actual": b["lineups"], "target": targets["lineups_per_day"]},
                {"metric": "Joinings / month", "actual": b["joined"], "target": targets["joinings_per_month"]},
            ]})
            for m in cards[-1]["metrics"]:
                m["pct"] = round(100 * m["actual"] / m["target"]) if m["target"] else 0
        return {"targets": cards, "score_formula": "calls*1 + connected*2 + lineups*5 + selected*8 + joined*15 - missed*2"}
    return {"leaderboard": board, "score_formula": "calls*1 + connected*2 + lineups*5 + selected*8 + joined*15 - missed*2"}


# ======================= ACTION REQUIRED =======================
async def action_required(db, actor, perms):
    scope_q = await scope_service.leads_filter(db, actor, perms, None)
    ids = await scope_service.recruiter_scope_ids(db, actor, perms)
    fu_scope = {} if ids is None else {"recruiter_id": {"$in": ids}}
    now = now_utc()

    async def sample(coll, q, limit=5, name_field="name"):
        rows = await db[coll].find(q, {"_id": 0}).limit(limit).to_list(limit)
        return rows

    stale_cut = now - timedelta(days=15)
    cards = []
    # 1 overdue followups
    q = {**fu_scope, "status": "pending", "due_at": {"$lt": now}}
    cards.append({"key": "overdue_followups", "title": "Overdue Follow-ups", "tone": "red",
                  "count": await db.followups.count_documents(q), "items": await sample("followups", q, name_field="lead_name")})
    # 2 active no followup
    q = {**scope_q, "status": {"$in": ACTIVE_STATUSES}, "$or": [{"next_followup_at": None}, {"next_followup_at": {"$exists": False}}]}
    cards.append({"key": "no_followup", "title": "Active Leads — No Follow-up", "tone": "amber",
                  "count": await db.leads.count_documents(q), "items": await sample("leads", q)})
    # 3 never called
    q = {**scope_q, "call_count": {"$lte": 0}, "status": {"$in": ACTIVE_STATUSES}}
    cards.append({"key": "never_called", "title": "Never-Called Leads", "tone": "amber",
                  "count": await db.leads.count_documents(q), "items": await sample("leads", q)})
    # 4 unassigned
    q = {"$or": [{"owner_id": None}, {"owner_id": ""}]}
    cards.append({"key": "unassigned", "title": "Unassigned Leads", "tone": "red",
                  "count": await db.leads.count_documents(q), "items": await sample("leads", q)})
    # 5 stale
    q = {**scope_q, "status": {"$in": ACTIVE_STATUSES}, "created_at": {"$lt": stale_cut}}
    cards.append({"key": "stale", "title": "Stale Leads (15+ days)", "tone": "amber",
                  "count": await db.leads.count_documents(q), "items": await sample("leads", q)})
    # 6 selected no joining date
    q = {**scope_q, "status": "selected", "$or": [{"expected_joining_at": None}, {"expected_joining_at": {"$exists": False}}]}
    cards.append({"key": "selected_no_joining", "title": "Selected — No Joining Date", "tone": "amber",
                  "count": await db.leads.count_documents(q), "items": await sample("leads", q)})
    # 7 unconfirmed interviews tomorrow
    tom = now.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)
    q = {**fu_scope, "status": "scheduled", "scheduled_at": {"$gte": tom, "$lt": tom + timedelta(days=1)}}
    cards.append({"key": "unconfirmed_interviews", "title": "Unconfirmed Interviews Tomorrow", "tone": "amber",
                  "count": await db.interviews.count_documents(q), "items": await sample("interviews", q, name_field="lead_name")})
    # 8 unconfirmed joinings
    q = {**fu_scope, "status": {"$in": ["pending", "confirmed"]}, "confirmation": {"$ne": "confirmed"}}
    cards.append({"key": "unconfirmed_joinings", "title": "Unconfirmed Joinings", "tone": "amber",
                  "count": await db.joinings.count_documents(q), "items": await sample("joinings", q, name_field="lead_name")})
    # 9 recruiters below call target
    org = await db.organization_settings.find_one({"key": "org"}) or {}
    target = (org.get("targets") or {}).get("calls_per_day", 50)
    umap = await _recruiter_ids(db, actor, perms)
    today = now.replace(hour=0, minute=0, second=0, microsecond=0)
    below = []
    for uid, name in umap.items():
        c = await db.calls.count_documents({"recruiter_id": uid, "created_at": {"$gte": today}})
        if c < target:
            below.append({"name": name, "calls": c, "target": target})
    cards.append({"key": "below_target", "title": "Recruiters Below Call Target", "tone": "red",
                  "count": len(below), "items": below[:5]})
    return {"cards": cards}


# ======================= INTEGRATIONS =======================
INTEGRATIONS = [
    {"key": "whatsapp", "name": "WhatsApp Business API", "fields": ["phone_number_id", "waba_id", "access_token", "webhook_verify_token", "app_secret"]},
    {"key": "instagram", "name": "Instagram / Meta Leads", "fields": ["page_id", "access_token", "webhook_verify_token", "app_secret"]},
    {"key": "facebook", "name": "Facebook Lead Ads", "fields": ["page_id", "access_token", "webhook_verify_token", "app_secret"]},
    {"key": "google_forms", "name": "Google Forms", "fields": ["shared_secret"]},
    {"key": "mobile_calling", "name": "Mobile Calling", "coming_soon": True, "fields": []},
]


async def list_integrations(db):
    out = []
    for meta in INTEGRATIONS:
        rec = await db.integrations.find_one({"key": meta["key"]})
        masked = {}
        if rec:
            for f in meta["fields"]:
                masked[f] = _mask(rec.get("_dec", {}).get(f)) if False else ("set" if rec.get("fields", {}).get(f) else "")
        out.append({**meta, "connected": bool(rec and rec.get("connected")),
                    "status": (rec or {}).get("status", "not_connected"),
                    "last_sync": (rec or {}).get("last_sync"),
                    "last_error": (rec or {}).get("last_error"),
                    "masked": {f: rec["fields"][f] for f in (rec.get("fields", {}) if rec else {})} if rec else {}})
    return out


async def save_integration(db, actor, key, values):
    meta = next((m for m in INTEGRATIONS if m["key"] == key), None)
    if not meta or meta.get("coming_soon"):
        raise AppError("bad_request", "Unsupported integration", 400)
    enc, masked = {}, {}
    for f in meta["fields"]:
        v = (values or {}).get(f)
        if v:
            enc[f] = _encrypt(v)
            masked[f] = _mask(v)
    await db.integrations.update_one({"key": key}, {"$set": {
        "key": key, "secrets_enc": enc, "fields": masked, "connected": True,
        "status": "connected", "updated_at": now_utc(), "updated_by": actor["id"],
    }}, upsert=True)
    return {"ok": True}


async def test_integration(db, key):
    rec = await db.integrations.find_one({"key": key})
    if not rec or not rec.get("connected"):
        raise AppError("not_connected", "Integration not configured", 400)
    await db.integrations.update_one({"key": key}, {"$set": {"last_sync": now_utc(), "status": "connected", "last_error": None}})
    return {"ok": True, "message": "Credentials present and stored securely (safe check passed)"}


async def disconnect_integration(db, key):
    await db.integrations.update_one({"key": key}, {"$set": {"connected": False, "status": "not_connected", "secrets_enc": {}, "fields": {}}})
    return {"disconnected": True}


# ======================= SETTINGS: org + lists =======================
DEFAULT_ORG = {
    "key": "org", "agency_name": "OAKsphere Recruitment",
    "targets": {"calls_per_day": 50, "connected_per_day": 20, "lineups_per_day": 3,
                "joinings_per_month": 5, "followup_escalation_hours": 24},
    "lead_sources": ["manual", "referral", "meta", "google_form", "careers", "whatsapp"],
    "priorities": ["hot", "high", "medium", "low", "cold"],
    "lead_statuses": list(STATUS_LABELS.keys()),
}


async def get_org_settings(db):
    rec = await db.organization_settings.find_one({"key": "org"}, {"_id": 0})
    if not rec:
        await db.organization_settings.insert_one(dict(DEFAULT_ORG))
        return dict(DEFAULT_ORG)
    for k, v in DEFAULT_ORG.items():
        rec.setdefault(k, v)
    return rec


async def update_org_settings(db, actor, data):
    cur = await get_org_settings(db)
    before = {k: cur.get(k) for k in data}
    await db.organization_settings.update_one({"key": "org"}, {"$set": data}, upsert=True)
    return {"before": before, "after": data}


# ======================= TAGS =======================
async def list_tags(db):
    rows = await db.lead_tags.find({}, {"_id": 0}).sort("name", 1).to_list(500)
    for r in rows:
        r["lead_count"] = await db.leads.count_documents({"tags": r["name"]})
    return rows


async def save_tag(db, data, tag_id=None):
    if not data.get("name"):
        raise AppError("bad_request", "Tag name is required", 400)
    now = now_utc()
    if tag_id:
        old = await db.lead_tags.find_one({"id": tag_id})
        if not old:
            raise AppError("not_found", "Tag not found", 404)
        if data.get("name") and data["name"] != old["name"]:
            await db.leads.update_many({"tags": old["name"]}, {"$set": {"tags.$": data["name"]}})
        await db.lead_tags.update_one({"id": tag_id}, {"$set": {**data, "updated_at": now}})
        return clean(await db.lead_tags.find_one({"id": tag_id}))
    if await db.lead_tags.find_one({"name": data["name"]}):
        raise AppError("duplicate", "Tag already exists", 400)
    doc = {"id": new_id(), "name": data["name"], "color": data.get("color", "indigo"), "created_at": now}
    await db.lead_tags.insert_one(doc)
    return clean(doc)


async def delete_tag(db, tag_id):
    tag = await db.lead_tags.find_one({"id": tag_id})
    if not tag:
        raise AppError("not_found", "Tag not found", 404)
    await db.leads.update_many({"tags": tag["name"]}, {"$pull": {"tags": tag["name"]}})
    await db.lead_tags.delete_one({"id": tag_id})
    return {"deleted": True}


# ======================= GLOBAL SEARCH =======================
async def global_search(db, actor, perms, q):
    if not q or len(q.strip()) < 2:
        return {"results": []}
    scope_q = await scope_service.leads_filter(db, actor, perms, None)
    norm = rs._normalize_phone(q)
    ors = [{"name": {"$regex": q, "$options": "i"}}, {"email": {"$regex": q, "$options": "i"}},
           {"lead_code": {"$regex": q, "$options": "i"}}]
    if norm:
        ors.append({"phone_normalized": {"$regex": norm}})
    query = {**scope_q, "$or": ors}
    rows = await db.leads.find(query, {"_id": 0, "id": 1, "name": 1, "phone": 1, "status": 1,
        "owner_id": 1, "job": 1, "client": 1, "lead_code": 1, "phone_normalized": 1}).limit(20).to_list(20)
    rows.sort(key=lambda r: 0 if r.get("phone_normalized") == norm else 1)
    return {"results": rows}


# ======================= IMPORT =======================
IMPORT_FIELDS = ["name", "phone", "alt_phone", "email", "city", "age", "gender", "qualification",
                 "experience", "current_salary", "expected_salary", "notice_period", "source",
                 "priority", "client", "job", "notes"]


def import_template_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([f.replace("_", " ").title() for f in IMPORT_FIELDS])
    ws.append(["John Doe", "9876543210", "", "john@example.com", "Pune", 28, "Male", "B.Com",
               "3 years", "30000", "40000", "30 days", "referral", "high", "Acme", "Sales Exec", "Warm lead"])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


def _parse_upload(filename, content):
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    elif name.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = [["" if c is None else str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise AppError("bad_request", "Only CSV and XLSX files are supported", 400)
    if not rows:
        raise AppError("bad_request", "File is empty", 400)
    return rows[0], rows[1:]


async def import_preview(db, actor, perms, filename, content):
    if len(content) > 5 * 1024 * 1024:
        raise AppError("bad_request", "File exceeds 5MB limit", 400)
    header, data_rows = _parse_upload(filename, content)
    phones_in_file = {}
    invalid = dupes_file = existing_dup = 0
    # tentative phone column guess
    pidx = next((i for i, h in enumerate(header) if "phone" in h.lower()), None)
    for r in data_rows:
        if pidx is not None and pidx < len(r):
            pn = rs._normalize_phone(r[pidx])
            if len(pn) != 10:
                invalid += 1
            if pn:
                phones_in_file[pn] = phones_in_file.get(pn, 0) + 1
    for pn, c in phones_in_file.items():
        if c > 1:
            dupes_file += c - 1
        if await db.leads.count_documents({"phone_normalized": pn}):
            existing_dup += 1
    batch_id = new_id()
    await db.import_batches.insert_one({
        "id": batch_id, "actor_id": actor["id"], "header": header, "rows": data_rows,
        "committed": False, "created_at": now_utc(),
        "expires_at": now_utc() + timedelta(hours=2),
    })
    return {"batch_id": batch_id, "columns": header, "row_count": len(data_rows),
            "invalid_phones": invalid, "duplicates_in_file": dupes_file,
            "existing_crm_duplicates": existing_dup,
            "preview": [dict(zip(header, r)) for r in data_rows[:5]],
            "supported_fields": IMPORT_FIELDS}


async def import_commit(db, actor, perms, batch_id, mapping, rules):
    batch = await db.import_batches.find_one({"id": batch_id})
    if not batch:
        raise AppError("not_found", "Import batch not found or expired", 404)
    if batch["committed"]:
        return batch.get("result", {"imported": 0, "note": "already committed"})
    if "phone" not in mapping.values():
        raise AppError("bad_request", "Phone column mapping is required", 400)
    header = batch["header"]
    col_index = {h: i for i, h in enumerate(header)}
    field_for = {col: field for col, field in mapping.items() if field and field != "not_mapped"}
    assign_mode = rules.get("assign", "unassigned")
    dup_rule = rules.get("duplicates", "flag")
    invalid_rule = rules.get("invalid", "flag")
    pool = []
    if assign_mode == "auto":
        sids = await scope_service.recruiter_scope_ids(db, actor, perms)
        pq = {"is_active": True, "role": {"$in": ["recruiter", "team_leader"]}}
        if sids is not None:
            pq["id"] = {"$in": sids}
        pool = [u["id"] for u in await db.users.find(pq, {"_id": 0, "id": 1}).to_list(200)]
    specific = rules.get("recruiter_id")

    res = {"imported": 0, "skipped_duplicates": 0, "skipped_invalid": 0,
           "flagged_duplicate": 0, "flagged_invalid": 0, "auto_distributed": 0}
    now = now_utc()
    docs = []
    rr = 0
    for r in batch["rows"]:
        def val(field):
            col = next((c for c, f in field_for.items() if f == field), None)
            if col is None or col_index.get(col) is None or col_index[col] >= len(r):
                return None
            return (r[col_index[col]] or "").strip() or None
        phone = val("phone")
        if not phone:
            res["skipped_invalid"] += 1; continue
        pn = rs._normalize_phone(phone)
        is_invalid = len(pn) != 10
        is_dup = bool(await db.leads.count_documents({"phone_normalized": pn}))
        if is_dup and dup_rule == "skip":
            res["skipped_duplicates"] += 1; continue
        if is_invalid and invalid_rule == "skip":
            res["skipped_invalid"] += 1; continue
        owner = None
        if assign_mode == "specific" and specific:
            owner = specific
        elif assign_mode == "auto" and pool:
            owner = pool[rr % len(pool)]; rr += 1; res["auto_distributed"] += 1
        age = val("age")
        doc = {"id": new_id(), "lead_code": rs._lead_code(), "name": val("name") or "Unknown",
               "phone": phone, "phone_normalized": pn, "alt_phone": val("alt_phone"),
               "email": (val("email") or "").lower() or None, "city": val("city"),
               "age": int(age) if (age and age.isdigit()) else None, "gender": (val("gender") or "").lower() or None,
               "qualification": val("qualification"), "experience": val("experience"),
               "current_salary": val("current_salary"), "expected_salary": val("expected_salary"),
               "notice_period": val("notice_period"), "source": val("source") or "import",
               "role_applied": val("job"), "client": val("client"), "job": val("job"),
               "priority": (val("priority") or "medium").lower(), "status": "new",
               "status_changed_at": now, "owner_id": owner, "team_id": None,
               "tags": [], "notes": val("notes"), "last_call_outcome": None, "call_count": 0,
               "last_activity_at": now, "next_followup_at": None, "closure_reason": None,
               "expected_joining_at": None, "invalid_phone_flag": is_invalid,
               "duplicate_flagged": is_dup, "import_batch_id": batch_id,
               "status_history": [{"status": "new", "at": now, "by": actor["id"]}],
               "created_at": now, "updated_at": now, "created_by": actor["id"], "updated_by": actor["id"]}
        docs.append(doc)
        res["imported"] += 1
        if is_dup:
            res["flagged_duplicate"] += 1
        if is_invalid:
            res["flagged_invalid"] += 1
    if docs:
        await db.leads.insert_many(docs)
    await db.import_batches.update_one({"id": batch_id}, {"$set": {"committed": True, "result": res}})
    return res
